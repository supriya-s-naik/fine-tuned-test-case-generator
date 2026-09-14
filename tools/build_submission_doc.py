from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Week5_Test_Case_Generator_Submission.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_table(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "D9D9D9")
        borders.append(tag)
    tbl_pr.append(borders)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.03
for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
styles["Title"].paragraph_format.space_after = Pt(10)
styles["Heading 1"].paragraph_format.space_before = Pt(11)
styles["Heading 1"].paragraph_format.space_after = Pt(4)
styles["Heading 2"].paragraph_format.space_before = Pt(7)
styles["Heading 2"].paragraph_format.space_after = Pt(2)

title = doc.add_paragraph(style="Title")
title.add_run("Fine Tuned Test Case Generator")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = subtitle.add_run("Week 5 Project  |  The Gen Academy  |  September 2026")
run.bold = True
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(55, 65, 81)

doc.add_paragraph(
    "I fine-tuned Qwen3 1.7B to generate structured software test cases from user stories and acceptance criteria. "
    "The evaluated adapter did not outperform the conversational base model overall: it scored 75.7% versus 76.2%. "
    "The experiment therefore shows that this small synthetic dataset does not yet justify fine-tuning for deployment."
)

doc.add_heading("Problem and use case", level=1)
doc.add_paragraph(
    "QA engineers repeatedly translate requirements into positive, negative, and boundary tests. General models can draft this material, "
    "but inconsistent structure and missing traceability make automation difficult. The target system accepts a user story plus numbered "
    "acceptance criteria and returns JSON test cases with id, title, type, preconditions, steps, expected result, and covered criterion IDs."
)

doc.add_heading("Dataset and validation design", level=1)
doc.add_paragraph(
    "The dataset contains 80 synthetic labelled examples across authentication, commerce, finance, healthcare, collaboration, analytics, "
    "travel, and learning products. Sixty examples from 12 scenario families were used for training. Twenty examples from four complete "
    "scenario families were held out. Keeping whole families out of training reduces leakage from closely related template variants."
)
doc.add_paragraph(
    "Each target contains one positive, one negative, and one boundary test. The covers field links every test to AC1, AC2, or AC3, "
    "making requirement coverage measurable rather than subjective."
)

doc.add_heading("Fine tuning approach", level=1)
settings = [
    ("Model", "Qwen/Qwen3-1.7B conversational checkpoint"),
    ("Method", "LoRA supervised fine-tuning in LLaMA Factory"),
    ("Learning rate", "1e-5"),
    ("Epochs", "2"),
    ("Batch", "1 per device with 4-step gradient accumulation"),
    ("LoRA", "Rank 8 and alpha 16"),
    ("Sequence and compute", "1024-token cutoff and FP16 on a T4 GPU"),
]
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(1.75)
table.columns[1].width = Inches(5.1)
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text = "Setting", "Evaluated run"
for cell in hdr:
    shade(cell, "243B53")
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
for idx, (key, value) in enumerate(settings, start=1):
    cells = table.add_row().cells
    cells[0].text, cells[1].text = key, value
    if idx % 2 == 0:
        for cell in cells:
            shade(cell, "F3F6FA")
    cells[0].paragraphs[0].runs[0].bold = True
for row in table.rows:
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margin(cell)
border_table(table)

doc.add_heading("Training iteration", level=2)
doc.add_paragraph(
    "An initial run on Qwen3 1.7B Base reached very low training loss but collapsed during autoregressive generation, repeating punctuation "
    "instead of producing JSON. I treated the smoke test as a release gate, rejected that model, switched to the conversational checkpoint, "
    "and reduced the learning rate. The second run produced coherent, traceable test cases on unseen appointment-booking stories."
)

doc.add_heading("Evaluation", level=1)
doc.add_paragraph(
    "The unchanged conversational base model and merged fine-tuned model received the same prompt, deterministic decoding, and 20 held-out "
    "examples. Exact-match accuracy would penalize valid wording differences, so I measured structural validity, requirement coverage, test "
    "breadth, and reference similarity."
)
metrics = [
    ("JSON validity", "100%", "100%", "0 pp"),
    ("Schema validity", "80%", "75%", "-5 pp"),
    ("AC coverage", "100%", "100%", "0 pp"),
    ("Test type coverage", "50%", "52%", "+2 pp"),
    ("ROUGE L", "51%", "52%", "+1 pp"),
    ("Overall mean", "76.2%", "75.7%", "-0.5 pp"),
]
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [3.0, 1.25, 1.35, 1.1]
for i, width in enumerate(widths):
    table.columns[i].width = Inches(width)
headers = ("Metric", "Base", "Fine tuned", "Delta")
for cell, value in zip(table.rows[0].cells, headers):
    cell.text = value
    shade(cell, "243B53")
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
for idx, values in enumerate(metrics, start=1):
    cells = table.add_row().cells
    for col, value in enumerate(values):
        cells[col].text = value
        cells[col].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.CENTER
    if idx % 2 == 0:
        for cell in cells:
            shade(cell, "F3F6FA")
    if values[0] == "Overall mean":
        for cell in cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True
for row in table.rows:
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margin(cell)
border_table(table)

doc.add_heading("Interpretation", level=1)
doc.add_paragraph(
    "Fine-tuning preserved perfect JSON validity and acceptance-criteria coverage and produced small gains in test-type coverage and ROUGE-L. "
    "Strict schema validity fell from 80% to 75%, leaving the composite score 0.5 percentage points below the base model. All five structural "
    "failures came from the held-out appointment-booking family, where expected_result was returned as a useful nonempty array rather than the required string."
)
doc.add_paragraph(
    "The content was coherent, but downstream integrations depend on the declared field types. I would not deploy this adapter over the base "
    "model on the current evidence. A further experiment should make every field type explicit, add diverse examples containing multi-part "
    "outcomes represented as single strings, and compare additional epochs against the same untouched validation set."
)

doc.add_heading("Notebook evidence", level=1)
doc.add_paragraph(
    "Before submitting, insert the screenshot from the final notebook cell directly below this paragraph. The screenshot should show "
    "EXPERIMENT COMPLETE, both model scores, the 20 held-out examples, and the -0.5 percentage-point result."
)
placeholder = doc.add_paragraph()
placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = placeholder.add_run("Insert final notebook screenshot here")
run.bold = True
run.font.color.rgb = RGBColor(80, 80, 80)
run.font.size = Pt(12)

doc.add_heading("Limitations", level=1)
doc.add_paragraph(
    "The dataset is synthetic and small. Structural metrics do not prove that every test is logically sufficient or free from unsupported "
    "assumptions. Production evaluation should use de-identified, QA-reviewed stories and human ratings for correctness, duplication, risk "
    "coverage, and actionability. Human review should remain in the workflow until performance is validated on real product requirements."
)

doc.save(OUT)
print(OUT)
